"""Render the companion paper MD draft to a Zenodo-ready DOCX with embedded figures.

Minimal Markdown subset renderer (headings, paragraphs, blockquotes, tables,
figure images, bold/italic/code inline) tuned for this manuscript. Run with the
training venv: .venv-train/bin/python experiments/build_docx.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
_stem = sys.argv[1] if len(sys.argv) > 1 else "COMPANION_PAPER_DRAFT_v0_2"
MD = ROOT / "docs" / f"{_stem}.md"
OUT = ROOT / "docs" / f"{_stem}.docx"

INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def add_runs(par, text):
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif piece.startswith("*") and piece.endswith("*"):
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


def main() -> int:
    lines = MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"; style.font.size = Pt(10.5)

    # strip YAML frontmatter
    if lines and lines[0].strip() == "---":
        end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
        lines = lines[end + 1:]

    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        # coalesce multi-line image markdown ![caption spanning lines](path)
        if s.startswith("!["):
            while not re.search(r"\]\([^)]*\)\s*$", s) and i + 1 < len(lines):
                i += 1
                s = s + " " + lines[i].strip()

        # figure: ![caption](path)
        m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
        if m:
            caption, rel = m.group(1), m.group(2)
            img = (MD.parent / rel).resolve()
            if img.exists():
                doc.add_picture(str(img), width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run(re.sub(r"\s+", " ", caption)); r.italic = True; r.font.size = Pt(9)
            i += 1
            continue

        # table block
        if s.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|:- "):
            header = [c.strip() for c in s.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]; cell.paragraphs[0].add_run(h).bold = True
            for row in rows:
                cells = t.add_row().cells
                for j, val in enumerate(row[:len(header)]):
                    add_runs(cells[j].paragraphs[0], val)
            doc.add_paragraph()
            continue

        if s.startswith("# "):
            doc.add_heading(s[2:], level=0)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=1)
        elif s.startswith("### "):
            doc.add_heading(s[4:], level=2)
        elif s.startswith("> "):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.4)
            r0 = p.add_run(""); add_runs(p, s[2:])
            for run in p.runs:
                run.italic = True; run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif s == "":
            pass
        else:
            add_runs(doc.add_paragraph(), s)
        i += 1

    doc.save(str(OUT))
    print(f"-> {OUT}  ({OUT.stat().st_size // 1024} KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
